"""Excel 导入：只有名称、没有编码的字典必须保留 + 告警，不得静默丢弃。

背景：v2.22.3 债清单记的是「Excel 侧 `if d.items` 会丢空字典」，实测**不准确** ——
导出侧（`import_export.py` `if d.items:` 的 else 分支）已能表达空字典，往返无损。
但沿这条线追下去挖出**真的静默数据丢失**：`data_standard_parser.py` 解析
「字典定义」表时，首列为空即 `continue`，字典**整个被丢弃且无任何告警**。

字典编码是**字段绑定**的连接键，但字典实体本身不依赖它 —— 丢弃整个实体是过度反应。
CSV 归档侧对同样情况有明确告警（「缺少字典编码」）并保留行，Excel 侧应对齐。
"""
import io
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(r"C:\Users\Administrator\Desktop\AI项目\空间质检")
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402
import pytest  # noqa: E402

from app.services.data_standard_parser import parse_standard_excel_v2  # noqa: E402

import tests.test_imp_exp_dict_roundtrip as T  # noqa: E402


def _xlsx_with_dict_rows(rows: list[list]) -> bytes:
    """构造一份含指定「字典定义」行的最小 Excel。

    「标准信息」是竖排 `键 | 值` 格式（解析器按 key 取 row[1]），
    「字典定义」是横排表头。
    """
    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = "标准信息"
    ws0.append(["字段名", "值"])
    ws0.append(["标准名称", "解析单测"])
    ws0.append(["标准编码", "PC-1"])
    ws0.append(["分类", "custom"])
    ws1 = wb.create_sheet("字典定义")
    ws1.append(["字典编码", "字典名称", "适用范围", "字典项编码", "字典项名称", "描述"])
    for row in rows:
        ws1.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse(rows: list[list]):
    """写出临时 xlsx 并解析（解析器收文件路径，不收字节流）。"""
    import tempfile
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as fh:
        fh.write(_xlsx_with_dict_rows(rows))
        path = fh.name
    try:
        return parse_standard_excel_v2(path)
    finally:
        Path(path).unlink(missing_ok=True)


class TestParserKeepsCodelessDictionary:
    """解析器层：只给名称的字典要活下来，并产生告警。"""

    def test_dict_without_code_kept_by_name(self):
        result = _parse([
            ["", "只有名称的字典", "standard", "", "", "说明"],
        ])
        names = [d.dict_name for d in result.dictionaries]
        assert "只有名称的字典" in names, f"无编码字典被丢弃：{result.dictionaries}"

    def test_dict_without_code_warns(self):
        result = _parse([
            ["", "只有名称的字典", "standard", "", "", ""],
        ])
        joined = " ".join(result.warnings)
        assert "只有名称的字典" in joined, result.warnings
        assert "编码" in joined, result.warnings

    def test_codeless_dict_keeps_its_items(self):
        """无编码但有字典项时，字典项不能跟着丢。"""
        result = _parse([
            ["", "无码字典", "standard", "A", "沥青", ""],
            ["", "无码字典", "standard", "B", "水泥", ""],
        ])
        d = next((x for x in result.dictionaries if x.dict_name == "无码字典"), None)
        assert d is not None, f"无编码字典被丢弃：{result.dictionaries}"
        assert sorted(d.items) == [("A", "沥青"), ("B", "水泥")], d.items

    def test_codeless_dicts_are_distinguished_by_name(self):
        """两个无编码字典名字不同 → 必须是两个实体，不能互相覆盖。"""
        result = _parse([
            ["", "字典甲", "standard", "", "", ""],
            ["", "字典乙", "standard", "", "", ""],
        ])
        names = sorted(d.dict_name for d in result.dictionaries)
        assert names == ["字典乙", "字典甲"], names

    def test_existing_coded_dict_behaviour_unchanged(self):
        """有编码的正常路径不受影响，且不产生新告警。"""
        result = _parse([
            ["MAT", "面材料", "standard", "A", "沥青", "说明"],
        ])
        assert [(d.dict_code, d.dict_name) for d in result.dictionaries] == [("MAT", "面材料")]
        assert not [w for w in result.warnings if "编码" in w], result.warnings

    def test_row_with_no_code_and_no_name_skipped(self):
        """编码和名称都空 → 跳过，否则会产生无名字典。

        注意要留一个非空单元格：全空行会被 `_rows_of` 在进循环前就滤掉，
        那样根本走不到这里的守卫，测试就成了「因错误的原因通过」。
        """
        result = _parse([
            ["", "", "", "", "", "只有描述没填编码和名称"],
        ])
        assert result.dictionaries == [], result.dictionaries

    def test_fully_blank_row_filtered_upstream(self):
        """整行全空：`_rows_of` 先滤掉，不产生任何字典、也不产生告警。"""
        result = _parse([
            [None, None, None, None, None, None],
        ])
        assert result.dictionaries == [], result.dictionaries
        assert not [w for w in result.warnings if "字典" in w], result.warnings


@pytest.mark.asyncio
async def test_excel_import_keeps_codeless_dictionary(client, db_session):
    """端点级：无编码字典经 Excel 导入后仍在库里，且有告警。"""
    raw, key = await T._key(db_session, "edl_1")
    std_id = await T._seed_rich_standard(db_session, key, "有码字典宿主")

    r = await client.post(f"/api/v1/data-standards-v2/{std_id}/export?format=excel",
                          headers={"X-API-Key": raw})
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["字典定义"]
    ws.append([None, "无码字典", "standard", None, None, "无编码的字典"])

    buf = io.BytesIO()
    wb.save(buf)
    r2 = await client.post("/api/v1/data-standards-v2/import/excel",
                           headers={"X-API-Key": raw},
                           data={"name": "无码字典回导", "mode": "update"},
                           files={"file": ("a.xlsx", buf.getvalue(), T.XLSX_CT)})
    assert r2.status_code == 201, r2.text
    snap = await T._snapshot(db_session, r2.json()["standard_id"])
    assert any(d[1] == "无码字典" for d in snap["dicts"]), snap["dicts"]
    assert any("无码字典" in w and "编码" in w for w in r2.json()["warnings"]), r2.json()["warnings"]
