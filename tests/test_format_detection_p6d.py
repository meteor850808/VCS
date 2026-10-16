"""P6d: format_detection.py → 100% 覆盖赛收官。

补齐 detect_format 置信度分支 / detect_geometry_type 各格式 / magic bytes
全分支 / CRS 提取异常路径 / schema 解析（DBF/KML/OBJ/CSV/XLSX）/
类型推断兜底。全部基于 tmp_path 真实文件，无 DB 依赖。
"""

import json
import sqlite3
import struct
import sys
from pathlib import Path

import pytest

from app.services import format_detection as fd
from app.services.format_detection import (
    FormatResult,
    detect_data_type,
    detect_format,
    detect_geometry_type,
    normalize_geometry_type,
    normalize_stored_geometry_type,
)


# ── 文件构造 helpers ─────────────────────────────────────────────────────────

def _write_shp(path: Path, shape_type: int) -> Path:
    """构造 SHP 主文件：前 4 字节文件码 9994 大端，shape_type 写入偏移 32。"""
    with open(path, "wb") as f:
        f.write(struct.pack(">i", 9994))
        f.write(b"\x00" * 28)
        f.write(struct.pack("<I", shape_type))
        f.write(b"\x00" * 64)
    return path


def _write_dbf(path: Path, fields: list[tuple[str, str]], extra=0) -> Path:
    """构造 DBF 文件头。fields: [(name, type_char), ...]，extra 追加填充字节。"""
    n = len(fields)
    header_len = 32 + 32 * n + 1
    buf = bytearray()
    buf += struct.pack("<B", 0x03)
    buf += b"\x00\x00\x00"
    buf += struct.pack("<I", 10)  # record count
    buf += struct.pack("<H", header_len)
    buf += struct.pack("<H", 1)  # record size
    buf += b"\x00" * 20
    for name, t in fields:
        name_b = name.encode("ascii")[:11].ljust(11, b"\x00")
        buf += name_b
        buf += t.encode("ascii")
        buf += b"\x00" * 4  # field data address
        buf += bytes([20])  # field length
        buf += bytes([0])  # decimal count
        buf += b"\x00" * 14
    buf += b"\x0D"
    buf += b"\x00" * extra
    path.write_bytes(bytes(buf))
    return path


def _write_geojson(path: Path, features: list[dict]) -> Path:
    path.write_text(json.dumps({"type": "FeatureCollection", "features": features}),
                    encoding="utf-8")
    return path


def _write_gpkg(path: Path, geom_type: str) -> Path:
    conn = sqlite3.connect(path)
    conn.execute("CREATE TABLE gpkg_geometry_columns (geometry_type_name TEXT)")
    conn.execute("INSERT INTO gpkg_geometry_columns VALUES (?)", (geom_type,))
    conn.commit()
    conn.close()
    return path


# ── 1. normalize_geometry_type 空值兜底 ───────────────────────────────────────

class TestNormalizeGeometryType:

    def test_empty_none(self):
        assert normalize_geometry_type("") is None
        assert normalize_geometry_type(None) is None

    def test_unknown_passthrough(self):
        assert normalize_geometry_type("BogusShape") == "BogusShape"

    def test_multipatch_3d(self):
        assert normalize_geometry_type("MultiPatch") == "3d"


# ── 2. detect_data_type 兜底分支 ──────────────────────────────────────────────

class TestDetectDataType:

    def test_mdb_format_spatial(self):
        assert detect_data_type("MDB") == "spatial"

    def test_mdb_filename_spatial(self):
        assert detect_data_type("test.mdb") == "spatial"

    def test_jet_db_content_fallback(self):
        """162：未知格式名 + Jet DB magic → spatial。"""
        assert detect_data_type("weird_thing", b"Standard Jet DB") == "spatial"

    def test_geometry_table_archive(self):
        assert detect_data_type("SHP") == "geometry"
        assert detect_data_type("x.shp") == "geometry"
        assert detect_data_type("CSV") == "table"
        assert detect_data_type("ZIP") == "archive"
        assert detect_data_type("a.tar.gz") == "archive"

    def test_unknown(self):
        assert detect_data_type("BOGUS") == "unknown"
        assert detect_data_type("noext") == "unknown"


# ── 3. detect_format 置信度 / 整体装配 ───────────────────────────────────────

class TestDetectFormat:

    def test_missing_file(self):
        r = detect_format("C:/definitely/not/exist.shp")
        assert r.format_name == "Unknown" and r.confidence == 0.0

    def test_confidence_06_ext_only(self, tmp_path):
        """195：扩展名命中但无 magic 匹配 → 0.6 + 幂等 layering。"""
        p = tmp_path / "data.csv"
        p.write_text("a,b\n1,2\n", encoding="utf-8")
        r = detect_format(str(p))
        assert r.format_name == "CSV"
        assert r.confidence == 0.6
        assert r.data_type == "table"
        assert r.geometry_type == "table"
        assert r.field_schema_json and r.field_schema_json[0]["name"] == "a"
        assert r.layer_count == 1

    def test_confidence_085_magic_mismatch(self, tmp_path):
        """191：magic 与扩展名不一致 → 0.85。"""
        p = tmp_path / "fake.kml"
        p.write_text('{"type":"FeatureCollection","features":[]}', encoding="utf-8")
        r = detect_format(str(p))
        assert r.format_name == "GeoJSON"
        assert r.confidence == 0.85

    def test_confidence_095_magic_match(self, tmp_path):
        """193：magic+ext 一致 → 0.95。"""
        p = tmp_path / "x.gpkg"
        _write_gpkg(p, "POINT")
        r = detect_format(str(p))
        assert r.format_name == "GPKG" and r.confidence == 0.95

    def test_confidence_01_unknown(self, tmp_path):
        """197：扩展名与 magic 均不识别 → 0.1。"""
        p = tmp_path / "blob.zzz"
        p.write_bytes(b"\x01\x02\x03\x04\x05")
        r = detect_format(str(p))
        assert r.format_name == "Unknown" and r.confidence == 0.1

    def test_directory_header_read_fails(self, tmp_path):
        """读目录 → _read_header_bytes 异常 → b'' → 0.1。"""
        d = tmp_path / "adir"
        d.mkdir()
        r = detect_format(str(d))
        assert r.confidence == 0.1


# ── 4. detect_geometry_type 各格式分支 ───────────────────────────────────────

class TestDetectGeometryType:

    def test_geojson_point(self, tmp_path):
        p = _write_geojson(tmp_path / "a.geojson", [
            {"type": "Feature", "geometry": {"type": "MultiPoint"},
             "properties": {}},
        ])
        assert detect_geometry_type(str(p), "GeoJSON") == "Point"

    def test_geojson_no_features(self, tmp_path):
        p = _write_geojson(tmp_path / "b.geojson", [])
        assert detect_geometry_type(str(p), "GeoJSON") is None

    def test_geojson_bad_json(self, tmp_path):
        p = tmp_path / "c.geojson"
        p.write_text("{not json", encoding="utf-8")
        assert detect_geometry_type(str(p), "GeoJSON") is None

    @pytest.mark.parametrize("shape_type,expected", [
        (1, "Point"), (3, "LineString"), (5, "Polygon"),
        (8, "Point"), (31, "3d"), (99, None),
    ])
    def test_shp_header_types(self, tmp_path, shape_type, expected):
        p = _write_shp(tmp_path / "t.shp", shape_type)
        assert detect_geometry_type(str(p), "SHP") == expected

    def test_shp_missing_main_file(self, tmp_path):
        """扩展名非 shp 且 .shp 邻居不存在 → None。"""
        p = tmp_path / "t.geojson"
        p.write_text("{}", encoding="utf-8")
        assert detect_geometry_type(str(p), "SHP") is None

    def test_shp_truncated_header(self, tmp_path):
        p = tmp_path / "t.shp"
        p.write_bytes(b"\x00" * 10)
        assert detect_geometry_type(str(p), "SHP") is None

    def test_kml_point(self, tmp_path):
        p = tmp_path / "a.kml"
        p.write_text(
            '<?xml version="1.0"?>'
            '<kml xmlns="http://www.opengis.net/kml/2.2">'
            "<Placemark><Point>1,2</Point></Placemark></kml>", encoding="utf-8")
        assert detect_geometry_type(str(p), "KML") == "Point"

    def test_kml_bad_xml(self, tmp_path):
        p = tmp_path / "bad.kml"
        p.write_text("not xml at all", encoding="utf-8")
        assert detect_geometry_type(str(p), "KML") is None

    def test_kml_no_geometry_tags(self, tmp_path):
        """299：XML 合法但无任何几何标签 → None。"""
        p = tmp_path / "plain.kml"
        p.write_text('<kml xmlns="http://www.opengis.net/kml/2.2">'
                     "<Document><name>x</name></Document></kml>", encoding="utf-8")
        assert detect_geometry_type(str(p), "KML") is None

    def test_gpkg_geometry_generic(self, tmp_path):
        """gpkg_geometry_columns 为 GEOMETRY → None。"""
        p = _write_gpkg(tmp_path / "mw.gpkg", "GEOMETRY")
        assert detect_geometry_type(str(p), "GPKG") is None

    def test_gpkg_point_z(self, tmp_path):
        """'POINT Z' → split → Point。"""
        p = _write_gpkg(tmp_path / "pz.gpkg", "POINT Z")
        assert detect_geometry_type(str(p), "GPKG") == "Point"

    def test_gpkg_bad_db(self, tmp_path):
        p = tmp_path / "bad.gpkg"
        p.write_bytes(b"\x00" * 8)
        assert detect_geometry_type(str(p), "GPKG") is None

    def test_obj_3d(self, tmp_path):
        assert detect_geometry_type("whatever.obj", "OBJ") == "3d"

    def test_table_formats(self, tmp_path):
        assert detect_geometry_type("a.csv", "CSV") == "table"
        assert detect_geometry_type("a.xlsx", "XLSX") == "table"

    def test_unhandled_fmt(self, tmp_path):
        assert detect_geometry_type("a.dxf", "DXF") is None

    def test_inner_helper_raises(self, tmp_path, monkeypatch):
        """239-240：子检测函数抛未捕获异常 → detect_geometry_type 兜底 None。"""

        def boom(path):
            raise RuntimeError("boom")

        monkeypatch.setattr(fd, "_detect_geom_geojson", boom)
        assert detect_geometry_type("x.geojson", "GeoJSON") is None


# ── 5. magic bytes 全分支 ─────────────────────────────────────────────────────

class TestDetectByMagic:

    def test_short_header(self):
        assert fd._detect_by_magic(b"\x00\x00") is None

    def test_shp_9994(self):
        assert fd._detect_by_magic(struct.pack(">i", 9994)) == "SHP"

    def test_shx_9993(self):
        assert fd._detect_by_magic(struct.pack(">i", 9993)) == "SHP"

    def test_dwg(self):
        assert fd._detect_by_magic(b"AC1027" + b"\x00" * 10) == "DWG"

    def test_tiff_le_be(self):
        assert fd._detect_by_magic(b"II\x2a\x00") == "TIFF"
        assert fd._detect_by_magic(b"MM\x00\x2a") == "TIFF"

    def test_gpkg_sqlite(self):
        assert fd._detect_by_magic(b"SQLite format 3\x00") == "GPKG"

    def test_mdb_jet(self):
        # header 恰 16 字节：b"Standard Jet DB" (15B) + \x00 → header[:22]==16B ✓
        assert fd._detect_by_magic(b"Standard Jet DB\x00") == "MDB"

    def test_geojson_text(self):
        assert fd._detect_by_magic(b'\xef\xbb\xbf{"type":"FeatureCollection"}') == "GeoJSON"

    def test_kml_xml(self):
        assert fd._detect_by_magic(b'<?xml version="1.0"?><kml xmlns="x">') == "KML"

    def test_gml_xml(self):
        assert fd._detect_by_magic(b"<?xml?><gml:FeatureCollection>") == "GML"

    def test_dxf_text(self):
        # 前导空白被 lstrip 吃掉，'  0\nSECTION' 须出现在中部
        assert fd._detect_by_magic(b"START\n  0\nSECTION\n  2\nHEADER") == "DXF"

    @pytest.mark.parametrize("line,expected", [
        (b"v 1 2 3\n", "OBJ"), (b"# comment\n", "OBJ"), (b"o box\n", "OBJ"),
    ])
    def test_obj_text_lines(self, line, expected):
        assert fd._detect_by_magic(line, ext=".obj") == expected

    def test_obj_wrong_ext(self):
        """OBJ 文本行但扩展名非 .obj → None。"""
        assert fd._detect_by_magic(b"v 1 2 3\n", ext=".txt") is None

    def test_no_match(self):
        assert fd._detect_by_magic(b"\x01\x02\x03\x04\x05\x06\x07\x08") is None


# ── 6. CRS 提取 ───────────────────────────────────────────────────────────────

class TestExtractCrs:

    def test_shp_prj(self, tmp_path):
        p = tmp_path / "a.shp"
        p.write_bytes(b"\x00" * 4)
        (tmp_path / "a.prj").write_text("GEOGCS[...]", encoding="utf-8")
        assert fd._extract_crs(str(p), "SHP") == "GEOGCS[...]"

    def test_shp_prj_read_error(self, tmp_path):
        """prj_exists 但不可读（目录）→ 静默 None。"""
        p = tmp_path / "a.shp"
        p.write_bytes(b"\x00" * 4)
        (tmp_path / "a.prj").mkdir()  # 目录: exists True, read_text 抛
        assert fd._extract_crs(str(p), "SHP") is None

    def test_shp_wo_prj(self, tmp_path):
        p = tmp_path / "b.shp"
        p.write_bytes(b"\x00" * 4)
        assert fd._extract_crs(str(p), "SHP") is None

    def test_geojson_crs_name(self, tmp_path):
        p = tmp_path / "c.geojson"
        p.write_text(json.dumps({"crs": {"type": "name",
                                         "properties": {"name": "EPSG:4326"}}}),
                     encoding="utf-8")
        assert fd._extract_crs(str(p), "GeoJSON") == "EPSG:4326"

    def test_geojson_crs_no_name(self, tmp_path):
        """crs 有定义但无 name → str(crs) 序列化。"""
        p = tmp_path / "d.geojson"
        p.write_text(json.dumps({"crs": {"type": "link"}}), encoding="utf-8")
        assert fd._extract_crs(str(p), "GeoJSON") == str({"type": "link"})

    def test_geojson_no_crs(self, tmp_path):
        p = tmp_path / "e.geojson"
        p.write_text("{}", encoding="utf-8")
        assert fd._extract_crs(str(p), "GeoJSON") is None

    def test_geojson_bad_json(self, tmp_path):
        p = tmp_path / "f.geojson"
        p.write_text("{broken", encoding="utf-8")
        assert fd._extract_crs(str(p), "GeoJSON") is None

    def test_kml_wgs84(self, tmp_path):
        p = tmp_path / "g.kml"
        p.write_text('<kml xmlns="..."><Placemark/></kml>', encoding="utf-8")
        assert fd._extract_crs(str(p), "KML") == "WGS84 (KML default)"

    def test_kml_without_kml_tag(self, tmp_path):
        p = tmp_path / "h.kml"
        p.write_text("<Whatever/>", encoding="utf-8")
        assert fd._extract_crs(str(p), "KML") is None

    def test_kml_unreadable(self, tmp_path):
        """读目录抛异常 → None。"""
        d = tmp_path / "i.kml"
        d.mkdir()
        assert fd._extract_crs(str(d), "KML") is None

    def test_obj_none(self, tmp_path):
        assert fd._extract_crs("x.obj", "OBJ") is None

    def test_other_fmt_none(self, tmp_path):
        assert fd._extract_crs("x.csv", "CSV") is None


# ── 7. schema 解析 ────────────────────────────────────────────────────────────

class TestParseSchema:

    def test_geojson_full(self, tmp_path):
        p = _write_geojson(tmp_path / "s.geojson", [
            {"type": "Feature",
             "geometry": {"type": "Point"},
             "properties": {"name": "A", "count": 3, "ratio": 1.5, "ok": True}},
            {"type": "Feature",
             "geometry": {"type": "LineString"},
             "properties": {"name": "B", "count": 4, "ratio": 2.5, "ok": False}},
        ])
        schema, count, layers = fd._parse_geojson_schema(str(p))
        assert count == 2 and layers == 1
        by = {s.name: s.field_type for s in schema}
        assert by == {"name": "string", "count": "integer",
                      "ratio": "float", "ok": "string"}

    def test_geojson_empty(self, tmp_path):
        p = _write_geojson(tmp_path / "e.geojson", [])
        schema, count, layers = fd._parse_geojson_schema(str(p))
        assert schema == [] and count == 0

    def test_geojson_no_properties_key(self, tmp_path):
        p = _write_geojson(tmp_path / "n.geojson", [
            {"type": "Feature", "geometry": {"type": "Point"}},
        ])
        schema, count, _ = fd._parse_geojson_schema(str(p))
        assert schema == [] and count == 1

    def test_shapefile_full_loop(self, tmp_path):
        """字段循环 + 类型映射 + 0x0D 终止。"""
        p = _write_dbf(tmp_path / "db.dbf",
                       [("FLD_C", "C"), ("FLD_N", "N"), ("FLD_I", "I"), ("FLD_X", "Z")])
        schema, count, layers = fd._parse_shapefile_schema(str(p))
        assert layers == 1 and count == 10
        by = {s.name: s.field_type for s in schema}
        assert by == {"FLD_C": "string", "FLD_N": "float",
                      "FLD_I": "integer", "FLD_X": "string"}

    def test_shapefile_no_dbf(self, tmp_path):
        assert fd._parse_shapefile_schema(str(tmp_path / "lonely.shp")) == ([], None, 0)

    def test_shapefile_short_header(self, tmp_path):
        """DBF 不足 32 字节 → 空 schema。"""
        p = tmp_path / "s.dbf"
        p.write_bytes(b"\x03" * 10)
        assert fd._parse_shapefile_schema(str(p)) == ([], None, 0)

    def test_shapefile_corrupt(self, tmp_path):
        """head_len 声明超出实际且无 0x0D 终止 → 字段循环 IndexError → 空。"""
        p = _write_dbf(tmp_path / "c.dbf", [("A", "C")])
        buf = bytearray(p.read_bytes())
        buf = buf[:-1]  # 去掉 0x0D 终止符，让循环读到 EOF → IndexError
        struct.pack_into("<H", buf, 8, 9999)
        p.write_bytes(bytes(buf))
        assert fd._parse_shapefile_schema(str(p)) == ([], None, 0)

    def test_shapefile_break_on_0x0d(self, tmp_path):
        """505：字段描述以 0x0D 开头 → break（跨 0x0D 的虚假 header_len）。"""
        p = _write_dbf(tmp_path / "d.dbf", [("A", "C")], extra=31)
        buf = bytearray(p.read_bytes())
        struct.pack_into("<H", buf, 8, 97)  # header_len-1 = 96 > 0x0D 位置 64
        p.write_bytes(bytes(buf))
        schema, count, layers = fd._parse_shapefile_schema(str(p))
        assert [s.name for s in schema] == ["A"]  # break 在附加字段处
        assert count == 10 and layers == 1

    def test_kml_simple_data(self, tmp_path):
        p = tmp_path / "k.kml"
        p.write_text(
            "<kml><Placemark><SimpleData name=\"city\">A</SimpleData>"
            "<SimpleData name=\"pop\">1</SimpleData></Placemark>"
            "<Placemark/></kml>", encoding="utf-8")
        schema, count, layers = fd._parse_kml_schema(str(p))
        assert {s.name for s in schema} == {"city", "pop"}
        assert all(s.field_type == "string" for s in schema)
        assert count == 1 and layers == 1  # <Placemark/> 不匹配 <Placemark>

    def test_kml_schema_definition(self, tmp_path):
        """无 SimpleData → 走 Schema 定义 + 类型映射。"""
        p = tmp_path / "k2.kml"
        p.write_text('<Schema><SimpleField name="id" type="int"/>'
                     '<SimpleField name="v" type="double"/></Schema>', encoding="utf-8")
        schema, count, _ = fd._parse_kml_schema(str(p))
        by = {s.name: s.field_type for s in schema}
        assert by == {"id": "integer", "v": "float"}
        assert count is None  # 无 Placemark

    def test_kml_unreadable(self, tmp_path):
        d = tmp_path / "kd.kml"
        d.mkdir()
        assert fd._parse_kml_schema(str(d)) == ([], None, 0)

    def test_obj_counts(self, tmp_path):
        p = tmp_path / "m.obj"
        p.write_text("v 0 0 0\nv 1 1 1\nf 1 2 3\nf 1 2 3\n"
                     "o box1\nv 2 2 2\no box2\n", encoding="utf-8")
        schema, count, layers = fd._parse_obj_schema(str(p))
        assert [s.name for s in schema] == ["vertex_x", "vertex_y", "vertex_z"]
        assert all(s.field_type == "float" for s in schema)
        assert count == 2 and layers == 2

    def test_obj_no_object(self, tmp_path):
        p = tmp_path / "m2.obj"
        p.write_text("v 0 0 0\n", encoding="utf-8")
        _, count, layers = fd._parse_obj_schema(str(p))
        assert count == 0 and layers == 1

    def test_obj_unreadable(self, tmp_path):
        d = tmp_path / "mo.obj"
        d.mkdir()
        assert fd._parse_obj_schema(str(d)) == ([], None, 0)

    def test_csv_full(self, tmp_path):
        p = tmp_path / "t.csv"
        p.write_text("id,name,price,,odd\n1,a,1.5,,\n2,b,2.5,,\n3,c,,x,y\n",
                     encoding="utf-8")
        schema, count, layers = fd._parse_csv_schema(str(p))
        assert count == 3 and layers == 1
        by = {s.name: (s.field_type, s.sample) for s in schema}
        assert by["id"] == ("integer", "1")
        assert by["name"] == ("string", "a")
        assert by["price"] == ("float", "1.5")
        assert by["col_4"] == ("string", "x")  # 空表头 → col_ 兜底
        assert by["odd"] == ("string", "y")

    def test_csv_empty(self, tmp_path):
        p = tmp_path / "e.csv"
        p.write_text("", encoding="utf-8")
        assert fd._parse_csv_schema(str(p)) == ([], None, 0)

    def test_csv_unreadable(self, tmp_path):
        d = tmp_path / "d.csv"
        d.mkdir()
        assert fd._parse_csv_schema(str(d)) == ([], None, 0)

    def test_xlsx_full(self, tmp_path):
        import openpyxl
        p = tmp_path / "t.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["name", "n", ""])
        ws.append(["a", 1, None])
        ws.append(["b", 2, None])
        wb.save(p)
        schema, count, layers = fd._parse_xlsx_schema(str(p))
        assert count == 2 and layers == 1
        by = {s.name: (s.field_type, s.sample) for s in schema}
        assert by["name"] == ("string", "a")
        assert by["n"] == ("integer", "1")
        assert by["col_3"] == ("string", None)  # 全空列 → sample None

    def test_xlsx_empty_sheet(self, tmp_path):
        import openpyxl
        p = tmp_path / "e.xlsx"
        wb = openpyxl.Workbook()
        wb.save(p)
        assert fd._parse_xlsx_schema(str(p)) == ([], None, 0)

    def test_xlsx_no_openpyxl(self, tmp_path, monkeypatch):
        monkeypatch.setitem(sys.modules, "openpyxl", None)
        assert fd._parse_xlsx_schema(str(tmp_path / "x.xlsx")) == ([], None, 0)

    def test_xlsx_corrupt(self, tmp_path):
        p = tmp_path / "c.xlsx"
        p.write_bytes(b"PK\x03\x04 not real zip")
        assert fd._parse_xlsx_schema(str(p)) == ([], None, 0)

    def test_xlsx_no_active_sheet(self, tmp_path, monkeypatch):
        """616：workbook 无 active sheet → 空。"""
        import openpyxl as _opx

        class FakeWB:
            active = None

        monkeypatch.setattr(_opx, "load_workbook", lambda *a, **k: FakeWB())
        assert fd._parse_xlsx_schema(str(tmp_path / "x.xlsx")) == ([], None, 0)

    def test_xlsx_blank_header(self, tmp_path):
        """632：表头为空白串 → col_ 命名兜底。"""
        import openpyxl
        p = tmp_path / "b.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["a", "  ", 1])
        wb.save(p)
        schema, _, _ = fd._parse_xlsx_schema(str(p))
        by = {s.name: s.field_type for s in schema}
        assert by["col_2"] == "string"  # "  " → strip → col_2

    def test_parse_schema_unhandled_fmt(self, tmp_path):
        assert fd._parse_schema("x.md", "BOGUS") == ([], None, 0)

    def test_parse_schema_geojson_bad(self, tmp_path):
        p = tmp_path / "b.geojson"
        p.write_text("{nope", encoding="utf-8")
        assert fd._parse_schema(str(p), "GeoJSON") == ([], None, 0)

    def test_parse_schema_kml_via_detect_format(self, tmp_path):
        """444：detect_format 全链路 → KML 分发。"""
        p = tmp_path / "a.kml"
        p.write_text('<kml xmlns="http://www.opengis.net/kml/2.2"><Placemark>'
                     '<SimpleData name="h">1</SimpleData>'
                     "<Point><coordinates>1,2</coordinates></Point>"
                     "</Placemark></kml>", encoding="utf-8")
        r = detect_format(str(p))
        assert r.format_name == "KML"
        assert r.geometry_type == "Point"
        assert {s.name for s in r.attribute_schema} == {"h"}
        assert r.feature_count == 1

    def test_parse_schema_obj_via_detect_format(self, tmp_path):
        """446：OBJ 分发。"""
        p = tmp_path / "m.obj"
        p.write_text("v 0 0 0\nv 1 1 1\nf 1 2 3\no box\n", encoding="utf-8")
        r = detect_format(str(p))
        assert r.format_name == "OBJ"
        assert r.geometry_type == "3d"
        assert [s.name for s in r.attribute_schema] == \
            ["vertex_x", "vertex_y", "vertex_z"]
        assert r.feature_count == 1  # face count
        assert r.layer_count == 1

    def test_parse_schema_xlsx_via_detect_format(self, tmp_path):
        """450：XLSX 分发。"""
        import openpyxl
        p = tmp_path / "t.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["name", "n"])
        ws.append(["a", 1])
        wb.save(p)
        r = detect_format(str(p))
        assert r.format_name == "XLSX"
        assert r.geometry_type == "table"
        by = {s.name: s.field_type for s in r.attribute_schema}
        assert by == {"name": "string", "n": "integer"}
        assert r.feature_count == 1


# ── 8. 类型推断 AND 兜底 map 分支 ────────────────────────────────────────────

class TestTypeHelpers:

    def test_infer_type_values(self):
        assert fd._infer_type(True) == "string"
        assert fd._infer_type(7) == "integer"
        assert fd._infer_type(1.5) == "float"
        assert fd._infer_type("x") == "string"
        assert fd._infer_type(None) == "string"

    def test_infer_csv_type_branches(self):
        assert fd._infer_csv_type(["", "  "]) == "string"  # 全空
        assert fd._infer_csv_type(["1", "2", "3"]) == "integer"
        assert fd._infer_csv_type(["1.5", "2.5"]) == "float"
        assert fd._infer_csv_type(["1", "2", "x"]) == "string"  # 66% < 80%

    def test_is_integer_float(self):
        assert fd._is_integer("42") is True
        assert fd._is_integer("4.2") is False
        assert fd._is_float("4.2") is True
        assert fd._is_float("abc") is False

    def test_dbf_type_fallback(self):
        assert fd._dbf_type("C") == "string"
        assert fd._dbf_type("F") == "float"
        assert fd._dbf_type("D") == "date"
        assert fd._dbf_type("?") == "string"  # 692 兜底

    def test_kml_type_fallback(self):
        assert fd._kml_type("int") == "integer"
        assert fd._kml_type("double") == "float"
        assert fd._kml_type("weird") == "string"  # 701 兜底

    def test_normalize_stored_full(self):
        assert normalize_stored_geometry_type("3D MultiPolygon") == "Polygon"
        assert normalize_stored_geometry_type("Polygon ZM") == "Polygon"
        assert normalize_stored_geometry_type("3D") == "3d"
        assert normalize_stored_geometry_type("2D") is None
        assert normalize_stored_geometry_type("MultiPatch") == "3d"
        assert normalize_stored_geometry_type("") is None
        assert normalize_stored_geometry_type("  ") is None


# ── 9. detect_format 完整装配（含 bbox 无关字段）──────────────────────────────

class TestDetectFormatAssembly:

    def test_shp_full_assembly(self, tmp_path):
        """完整 SMOKE：SHP + prj + dbf 装配层。"""
        _write_shp(tmp_path / "g.shp", 5)
        _write_dbf(tmp_path / "g.dbf", [("NAME", "C"), ("AREA", "F")])
        (tmp_path / "g.prj").write_text("PROJCS[...]", encoding="utf-8")
        r = detect_format(str(tmp_path / "g.shp"))
        assert r.format_name == "SHP"
        assert r.confidence == 0.95
        assert r.crs_info == "PROJCS[...]"
        assert r.feature_count == 10
        assert r.geometry_type == "Polygon"
        assert r.layer_count == 1
        assert r.magic_bytes == "0000270a"
        assert r.data_type == "geometry"
        assert r.attribute_schema and r.attribute_schema[0].name == "NAME"
        assert r.field_schema_json
        assert r.encoding

    def test_geojson_assembly_empty_schema(self, tmp_path):
        """无属性 → field_schema_json=None。"""
        p = _write_geojson(tmp_path / "n.geojson", [
            {"type": "Feature", "geometry": {"type": "Point"}, "properties": {}},
        ])
        r = detect_format(str(p))
        assert r.format_name == "GeoJSON"
        assert r.geometry_type == "Point"
        assert r.field_schema_json is None