# -*- coding: utf-8 -*-
"""P5h — app/api/v1/data_standard_v2/import_export.py 路由级 100% 覆盖。

覆盖：JSON 导入的 update/upsert/字典项/children 递归、Excel 导入（v2 模板/
旧扁平/坏模式/坏扩展名/.xls/空模板/409）、CSV 导入（create/update/409/
空文件/跳过非 CSV）、模板下载（excel/csv）、导出（excel/404/400/越权）、
auto-create-table（ok/404/无 schema/ValueError/越权）。

实现层函数直调：_import_version_snapshot(extra_diff 分支)、_content_disposition
（纯中文 fallback）。
"""
import io
import json
import uuid
import zipfile
from unittest.mock import AsyncMock, patch

import pytest
import pytest_asyncio
from openpyxl import Workbook
from sqlalchemy import select

from app.api.v1.data_standard_v2 import import_export as ie
from app.core.security import generate_api_key
from app.models.api_key import ApiKey
from app.models.data_standard import DataStandard
from app.models.dictionary_item import DictionaryItem
from app.models.source_data import StandardVersion
from app.models.standard_dictionary import StandardDictionary
from app.models.standard_field import StandardField
from app.models.standard_layer_group import StandardLayerGroup
from app.models.standard_table import StandardTable
from app.models.stored_file import StoredFile

pytestmark = pytest.mark.asyncio


@pytest_asyncio.fixture(scope="session", autouse=True)
async def _autouse_test_db(test_db):
    yield


async def _key(db_session, name):
    raw, hashed = generate_api_key()
    k = ApiKey(key_hash=hashed, name=name, is_active=True)
    db_session.add(k)
    await db_session.commit()
    return raw, k


# ── Excel 模板构造 ──────────────────────────────────────────────────────────

def _v2_xlsx(std_name="道路标准", code="RD_STD", category="custom", desc="道路标准"):
    """标准 v2 四 Sheet：1 根层 + 1 子层 + 1 表 + 2 字段 + 1 字典(1 项)。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "标准信息"
    ws.append(["字段名", "值"])
    ws.append(["标准名称", std_name])
    ws.append(["标准编码", code])
    ws.append(["分类", category])
    ws.append(["描述", desc])
    ws2 = wb.create_sheet("层组定义")
    ws2.append(["层组名", "层组编码", "父级编码", "描述"])
    ws2.append(["根层", "Root", "", ""])
    ws2.append(["子层", "Child", "Root", ""])
    ws3 = wb.create_sheet("数据表定义")
    ws3.append(["层组编码", "表名", "表编码", "表类型", "几何类型", "描述"])
    ws3.append(["Root", "面层", "Surf", "spatial", "Polygon", ""])
    ws4 = wb.create_sheet("字段定义")
    ws4.append(["表编码", "字段名", "中文名", "数据类型", "长度", "精度", "必填", "主键", "描述"])
    ws4.append(["Surf", "code", "编码", "string", "20", "", "是", "是", ""])
    ws4.append(["Surf", "name", "名称", "string", "50", "", "", "", ""])
    ws5 = wb.create_sheet("字典定义")
    ws5.append(["字典编码", "字典名称", "适用范围", "字典项编码", "字典项名称", "描述"])
    ws5.append(["MAT", "材料", "standard", "A", "沥青", ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _v2_xlsx_edge():
    """v2 怪模板：命中 fallback 组/表编码缺失/scope global/空字典项名等边角。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "标准信息"
    ws.append(["字段名", "值"])
    ws.append(["标准名称", "边角标准"])
    ws.append(["标准编码", "EDGE"])
    ws2 = wb.create_sheet("层组定义")
    ws2.append(["层组名", "层组编码", "父级编码", "描述"])
    ws2.append(["主组", "Main", "", ""])
    ws3 = wb.create_sheet("数据表定义")
    ws3.append(["层组编码", "表名", "表编码", "表类型", "几何类型", "描述"])
    ws3.append(["Main", "正常表", "Ok", "mystery", "Point", ""])   # 非法 type → attribute
    ws3.append(["NOPE", "孤儿表", "Orphan", "spatial", "", ""])     # 层组缺失 → fallback
    ws4 = wb.create_sheet("字段定义")
    ws4.append(["表编码", "字段名", "中文名", "数据类型", "长度", "精度", "必填", "主键", "描述"])
    ws4.append(["Ok", "f1", "", "string", "", "", "", "", ""])
    ws4.append(["GHOST", "幽灵字段", "", "string", "", "", "", "", ""])  # 表缺失 → skip
    ws5 = wb.create_sheet("字典定义")
    ws5.append(["字典编码", "字典名称", "适用范围", "字典项编码", "字典项名称", "描述"])
    ws5.append(["GL", "", "global", "A", "项目A", ""])
    ws5.append(["GL", "", "global", "B", "", ""])   # item_name 缺失 → 用 item_code
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _flat_xlsx_bytes():
    """旧扁平结构（非 4-Sheet）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "用地"
    ws.append(["字段名", "类型", "长度", "可空", "允许值", "描述"])
    ws.append(["name", "string", 10, "否", "", ""])
    ws.append(["area", "float", None, "是", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _empty_xlsx_bytes():
    """无有效工作表：表头无法识别 → 400。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["无意义", "内容"])
    ws.append(["a", "b"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _files(field, items):
    """构造 multipart files。"""
    return {field: items}


# ═══════════════════════════════════════════════════════════════════════════
# JSON 导入：update / upsert / code 匹配 / children 递归 / 字典项
# ═══════════════════════════════════════════════════════════════════════════

class TestImportJsonModes:

    async def test_json_update_by_name(self, client, db_session):
        raw, key = await _key(db_session, "ie_ju_a")
        h = {"X-API-Key": raw}
        p1 = {"name": "JSON标准A", "code": "JSA", "category": "custom",
              "layer_groups": [], "dictionaries": []}
        r1 = await client.post("/api/v1/data-standards-v2/import/json", headers=h, json=p1)
        assert r1.status_code == 201
        assert r1.json()["updated"] is False

        p2 = {"name": "JSON标准A", "code": "JSA", "category": "sf",
              "layer_groups": [], "dictionaries": []}
        r2 = await client.post(
            "/api/v1/data-standards-v2/import/json", headers=h, json=p2,
            params={"mode": "update"},
        )
        assert r2.status_code == 201
        body = r2.json()
        assert body["updated"] is True
        assert body["version"] >= 1
        assert body["name"] == "JSON标准A"

        # 版本快照已落库
        rows = (await db_session.execute(
            select(StandardVersion).where(StandardVersion.standard_id == body["standard_id"])
        )).scalars().all()
        assert len(rows) >= 1
        assert rows[0].change_log.startswith("JSON 导入更新")

    async def test_json_update_by_code_only(self, client, db_session):
        raw, key = await _key(db_session, "ie_ju_b")
        h = {"X-API-Key": raw}
        await client.post("/api/v1/data-standards-v2/import/json", headers=h,
                          json={"name": "原始名", "code": "BY_CODE"})
        r = await client.post(
            "/api/v1/data-standards-v2/import/json", headers=h,
            json={"name": "新名字", "code": "BY_CODE"},
            params={"mode": "update"},
        )
        assert r.status_code == 201
        assert r.json()["name"] == "新名字"
        assert r.json()["updated"] is True

    async def test_json_upsert_creates_when_missing(self, client, db_session):
        raw, _ = await _key(db_session, "ie_ju_c")
        r = await client.post(
            "/api/v1/data-standards-v2/import/json",
            headers={"X-API-Key": raw},
            json={"name": "UPSERT新", "code": "UPS-1"},
            params={"mode": "upsert"},
        )
        assert r.status_code == 201
        assert r.json()["updated"] is False

    async def test_json_update_not_found_404(self, client, db_session):
        raw, _ = await _key(db_session, "ie_ju_d")
        r = await client.post(
            "/api/v1/data-standards-v2/import/json",
            headers={"X-API-Key": raw},
            json={"name": "不存在的标准"},
            params={"mode": "update"},
        )
        assert r.status_code == 404

    async def test_json_bad_mode_400(self, client, db_session):
        raw, _ = await _key(db_session, "ie_ju_e")
        r = await client.post(
            "/api/v1/data-standards-v2/import/json",
            headers={"X-API-Key": raw},
            json={},
            params={"mode": "nope"},
        )
        assert r.status_code == 400
        assert "create" in r.json()["detail"]

    async def test_json_nested_children(self, client, db_session):
        raw, _ = await _key(db_session, "ie_ju_f")
        payload = {
            "name": "嵌套标准",
            "layer_groups": [
                {
                    "group_name": "根",
                    "group_code": "R",
                    "tables": [{"table_name": "表根", "fields": [{"field_name": "f0"}]}],
                    "children": [
                        {
                            "group_name": "子",
                            "group_code": "C",
                            "tables": [{"table_name": "表子", "fields": [{"field_name": "f1"}, {"field_name": "f2"}]}],
                            "children": [],
                        }
                    ],
                }
            ],
            "dictionaries": [],
        }
        r = await client.post("/api/v1/data-standards-v2/import/json",
                              headers={"X-API-Key": raw}, json=payload)
        assert r.status_code == 201
        body = r.json()
        assert body["tables_count"] == 2
        assert body["fields_count"] == 3
        std_id = body["standard_id"]

        layers = (await db_session.execute(
            select(StandardLayerGroup).where(StandardLayerGroup.standard_id == std_id)
        )).scalars().all()
        assert len(layers) == 2
        child = next(l for l in layers if l.group_name == "子")
        assert child.parent_id is not None

    async def test_json_dict_with_items(self, client, db_session):
        raw, _ = await _key(db_session, "ie_ju_g")
        payload = {
            "name": "字典标准",
            "dictionaries": [
                {
                    "dict_name": "一级分类",
                    "dict_code": "LV1",
                    "items": [
                        {"item_code": "A", "item_name": "甲", "parent_code": None},
                        {"item_code": "B", "item_name": "乙", "parent_code": "A"},
                    ],
                }
            ],
        }
        r = await client.post("/api/v1/data-standards-v2/import/json",
                              headers={"X-API-Key": raw}, json=payload)
        assert r.status_code == 201
        assert r.json()["dictionaries_count"] == 1
        std_id = r.json()["standard_id"]
        items = (await db_session.execute(
            select(DictionaryItem).join(StandardDictionary)
            .where(StandardDictionary.standard_id == std_id)
        )).scalars().all()
        assert len(items) == 2
        assert any(i.parent_code == "A" for i in items)


# ═══════════════════════════════════════════════════════════════════════════
# Excel 导入
# ═══════════════════════════════════════════════════════════════════════════

class TestImportExcel:

    async def test_excel_v2_create(self, client, db_session):
        raw, _ = await _key(db_session, "ie_e1_a")
        r = await client.post(
            "/api/v1/data-standards-v2/import/excel",
            headers={"X-API-Key": raw},
            data={"name": "道路标准", "mode": "create"},
            files={"file": ("road.xlsx", _v2_xlsx(), XLSX_CT)},
        )
        assert r.status_code == 201
        body = r.json()
        assert body["layer_groups_count"] == 2
        assert body["tables_count"] == 1
        assert body["fields_count"] == 2
        assert body["dictionaries_count"] == 1
        # 子层 parent 已被挂到根层
        std_id = body["standard_id"]
        layers = (await db_session.execute(
            select(StandardLayerGroup).where(StandardLayerGroup.standard_id == std_id)
        )).scalars().all()
        child = next(l for l in layers if l.group_name == "子层")
        assert child.parent_id is not None
        table = (await db_session.execute(
            select(StandardTable).where(StandardTable.standard_id == std_id)
        )).scalars().one()
        assert table.geometry_type == "Polygon"

    async def test_excel_v2_update(self, client, db_session):
        raw, _ = await _key(db_session, "ie_e1_b")
        h = {"X-API-Key": raw}
        r1 = await client.post("/api/v1/data-standards-v2/import/excel", headers=h,
                               data={"name": "道路标准"}, files={"file": ("a.xlsx", _v2_xlsx(), XLSX_CT)})
        assert r1.status_code == 201
        r2 = await client.post("/api/v1/data-standards-v2/import/excel", headers=h,
                               data={"name": "道路标准", "mode": "update"},
                               files={"file": ("b.xlsx", _v2_xlsx(), XLSX_CT)})
        assert r2.status_code == 201
        assert r2.json()["updated"] is True
        assert r2.json()["version"] >= 1

    async def test_excel_v2_edge_branches(self, client, db_session):
        raw, _ = await _key(db_session, "ie_e1_c")
        r = await client.post(
            "/api/v1/data-standards-v2/import/excel",
            headers={"X-API-Key": raw},
            data={"name": "边角标准", "mode": "create"},
            files={"file": ("edge.xlsx", _v2_xlsx_edge(), XLSX_CT)},
        )
        assert r.status_code == 201
        body = r.json()
        assert body["tables_count"] == 2
        assert body["fields_count"] == 1
        assert body["dictionaries_count"] == 1
        hits = [w for w in body["warnings"] if "未分组" in w]
        assert len(hits) == 1
        ghost = [w for w in body["warnings"] if "已跳过" in w]
        assert len(ghost) == 1
        assert "字典项" in body["message"]

        std_id = body["standard_id"]
        table = (await db_session.execute(
            select(StandardTable).where(StandardTable.standard_id == std_id,
                                        StandardTable.table_name == "正常表")
        )).scalars().one()
        assert table.table_type == "attribute"  # mystery → attribute
        orphan_table = (await db_session.execute(
            select(StandardTable).where(StandardTable.standard_id == std_id,
                                        StandardTable.table_name == "孤儿表")
        )).scalars().one()
        orphan_layer = (await db_session.execute(
            select(StandardLayerGroup).where(StandardLayerGroup.id == orphan_table.layer_id)
        )).scalars().one()
        assert orphan_layer.group_name == "未分组"
        # scope=global 字典 + 空 item_name 用 code
        d = (await db_session.execute(
            select(StandardDictionary).where(StandardDictionary.standard_id == std_id)
        )).scalars().one()
        assert d.scope == "global"
        items = (await db_session.execute(
            select(DictionaryItem).where(DictionaryItem.dict_id == d.id)
        )).scalars().all()
        assert len(items) == 2
        assert any(i.item_name == "B" for i in items)  # item_name 回退 item_code

    async def test_excel_v2_409_duplicate(self, client, db_session):
        raw, _ = await _key(db_session, "ie_e1_d")
        h = {"X-API-Key": raw}
        await client.post("/api/v1/data-standards-v2/import/excel", headers=h,
                          data={"name": "撞名标准"},
                          files={"file": ("a.xlsx", _v2_xlsx("撞名标准", "DUPE"), XLSX_CT)})
        r2 = await client.post("/api/v1/data-standards-v2/import/excel", headers=h,
                               data={"name": "撞名标准"},
                               files={"file": ("b.xlsx", _v2_xlsx("撞名标准", "DUPE"), XLSX_CT)})
        assert r2.status_code == 409

    async def test_excel_flat_create(self, client, db_session):
        raw, _ = await _key(db_session, "ie_e2_a")
        r = await client.post(
            "/api/v1/data-standards-v2/import/excel",
            headers={"X-API-Key": raw},
            data={"name": "旧标准", "mode": "create"},
            files={"file": ("old.xlsx", _flat_xlsx_bytes(), XLSX_CT)},
        )
        assert r.status_code == 201
        body = r.json()
        assert body["tables_count"] == 1
        assert body["fields_count"] == 2
        assert body["layer_groups_count"] == 1
        std = (await db_session.get(DataStandard, body["standard_id"]))
        assert std.file_name == "old.xlsx"

    async def test_excel_flat_update(self, client, db_session):
        raw, _ = await _key(db_session, "ie_e2_b")
        h = {"X-API-Key": raw}
        await client.post("/api/v1/data-standards-v2/import/excel", headers=h,
                          data={"name": "旧更新标准"},
                          files={"file": ("a.xlsx", _flat_xlsx_bytes(), XLSX_CT)})
        r2 = await client.post("/api/v1/data-standards-v2/import/excel", headers=h,
                               data={"name": "旧更新标准", "mode": "update"},
                               files={"file": ("b.xlsx", _flat_xlsx_bytes(), XLSX_CT)})
        assert r2.status_code == 201
        assert r2.json()["updated"] is True
        assert r2.json()["version"] >= 1

    async def test_excel_flat_409(self, client, db_session):
        raw, _ = await _key(db_session, "ie_e2_c")
        h = {"X-API-Key": raw}
        for i in range(2):
            r = await client.post("/api/v1/data-standards-v2/import/excel", headers=h,
                                  data={"name": "旧撞名"},
                                  files={"file": (f"f{i}.xlsx", _flat_xlsx_bytes(), XLSX_CT)})
        assert r.status_code == 409

    async def test_excel_bad_mode(self, client, db_session):
        raw, _ = await _key(db_session, "ie_e3_a")
        r = await client.post("/api/v1/data-standards-v2/import/excel",
                              headers={"X-API-Key": raw},
                              data={"name": "x", "mode": "nope"},
                              files={"file": ("a.xlsx", _v2_xlsx(), XLSX_CT)})
        assert r.status_code == 400

    async def test_excel_bad_extension(self, client, db_session):
        raw, _ = await _key(db_session, "ie_e3_b")
        r = await client.post("/api/v1/data-standards-v2/import/excel",
                              headers={"X-API-Key": raw},
                              data={"name": "x"},
                              files={"file": ("a.txt", b"x", "text/plain")})
        assert r.status_code == 400

    async def test_excel_xls_rejected(self, client, db_session):
        raw, _ = await _key(db_session, "ie_e3_c")
        r = await client.post("/api/v1/data-standards-v2/import/excel",
                              headers={"X-API-Key": raw},
                              data={"name": "x"},
                              files={"file": ("a.xls", b"", "application/vnd.ms-excel")})
        assert r.status_code == 400
        assert "xlsx" in r.json()["detail"]

    async def test_excel_no_valid_sheets(self, client, db_session):
        raw, _ = await _key(db_session, "ie_e3_d")
        r = await client.post("/api/v1/data-standards-v2/import/excel",
                              headers={"X-API-Key": raw},
                              data={"name": "x"},
                              files={"file": ("empty.xlsx", _empty_xlsx_bytes(), XLSX_CT)})
        assert r.status_code == 400
        assert "有效的工作表" in r.json()["detail"]


# ═══════════════════════════════════════════════════════════════════════════
# CSV 导入
# ═══════════════════════════════════════════════════════════════════════════

class TestImportCsv:

    async def test_csv_create_with_skip(self, client, db_session):
        raw, _ = await _key(db_session, "ie_c1_a")
        csv_bytes = "col1,col2,col3\na,b,c\n1,2,3\n".encode("utf-8")
        r = await client.post(
            "/api/v1/data-standards-v2/import/csv",
            headers={"X-API-Key": raw},
            data={"name": "CSV标准", "mode": "create"},
            files=[("files", ("t1.csv", csv_bytes, "text/csv")),
                   ("files", ("bad.txt", b"abc", "text/plain"))],
        )
        assert r.status_code == 201
        body = r.json()
        assert body["tables_count"] == 1
        assert body["fields_count"] == 3
        assert len([w for w in body["warnings"] if "非 CSV" in w]) == 1
        std = (await db_session.get(DataStandard, body["standard_id"]))
        layer = (await db_session.execute(
            select(StandardLayerGroup).where(StandardLayerGroup.standard_id == std.id)
        )).scalars().one()
        assert "导入层组" in layer.group_name
        table = (await db_session.execute(
            select(StandardTable).where(StandardTable.standard_id == std.id)
        )).scalars().one()
        assert table.table_type == "attribute"

    async def test_csv_update(self, client, db_session):
        raw, _ = await _key(db_session, "ie_c1_b")
        h = {"X-API-Key": raw}
        csv_bytes = b"x1,x2\n"
        await client.post("/api/v1/data-standards-v2/import/csv", headers=h,
                          data={"name": "CSV更新标准"},
                          files=[("files", ("a.csv", csv_bytes, "text/csv"))])
        r2 = await client.post("/api/v1/data-standards-v2/import/csv", headers=h,
                               data={"name": "CSV更新标准", "mode": "update"},
                               files=[("files", ("b.csv", b"y1,y2,y3\n", "text/csv"))])
        assert r2.status_code == 201
        body = r2.json()
        assert body["updated"] is True
        assert body["version"] >= 1

    async def test_csv_409(self, client, db_session):
        raw, _ = await _key(db_session, "ie_c1_c")
        h = {"X-API-Key": raw}
        csv_bytes = b"x1\n"
        for i in range(2):
            r = await client.post("/api/v1/data-standards-v2/import/csv", headers=h,
                                  data={"name": "CSV撞名"},
                                  files=[("files", (f"f{i}.csv", csv_bytes, "text/csv"))])
        assert r.status_code == 409

    async def test_csv_bad_mode(self, client, db_session):
        raw, _ = await _key(db_session, "ie_c1_d")
        r = await client.post("/api/v1/data-standards-v2/import/csv",
                              headers={"X-API-Key": raw},
                              data={"name": "x", "mode": "nope"},
                              files=[("files", ("a.csv", b"x\n", "text/csv"))])
        assert r.status_code == 400

    async def test_csv_empty_and_only_skips(self, client, db_session):
        raw, _ = await _key(db_session, "ie_c1_e")
        r = await client.post(
            "/api/v1/data-standards-v2/import/csv",
            headers={"X-API-Key": raw},
            data={"name": "CSV空文件"},
            files=[("files", ("empty.csv", b"", "text/csv")),
                   ("files", ("note.txt", b"abc", "text/plain"))],
        )
        assert r.status_code == 201
        body = r.json()
        assert body["tables_count"] == 0
        assert body["fields_count"] == 0
        assert any("空文件" in w for w in body["warnings"])
        assert any("非 CSV" in w for w in body["warnings"])


# ═══════════════════════════════════════════════════════════════════════════
# 模板下载 + 导出
# ═══════════════════════════════════════════════════════════════════════════

class TestTemplatesAndExport:

    async def test_download_excel_template(self, client):
        r = await client.get("/api/v1/data-standards-v2/template/excel")
        assert r.status_code == 200
        assert r.content[:2] == b"PK"
        assert "standard_template.xlsx" in r.headers["content-disposition"]

    async def test_download_csv_template(self, client):
        """模板 = manifest + 6 个角色 CSV（v2.22.3 起不再是一表一文件的示例包）。"""
        r = await client.get("/api/v1/data-standards-v2/template/csv")
        assert r.status_code == 200
        assert r.content[:4] == b"PK\x03\x04"
        with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
            names = zf.namelist()
            manifest = json.loads(zf.read("_manifest.json").decode("utf-8"))
        assert manifest["format"] == "vcs-standard-csv"
        assert manifest["version"] == 1
        for filename in manifest["roles"].values():
            assert filename in names, f"模板缺少角色文件: {filename}"

    async def test_export_excel_full(self, client, db_session):
        raw, _ = await _key(db_session, "ie_x1_a")
        r = await client.post(
            "/api/v1/data-standards-v2/import/json",
            headers={"X-API-Key": raw},
            json={
                "name": "导出标准",
                "code": "EXP-1",
                "layer_groups": [{
                    "group_name": "根", "group_code": "R",
                    "tables": [{"table_name": "T", "table_code": "TC",
                                "geometry_type": "Point",
                                "fields": [{"field_name": "f1", "data_type": "string"}]}],
                    "children": [],
                }],
                "dictionaries": [
                    {"dict_name": "D1", "dict_code": "DX",
                     "items": [{"item_code": "a", "item_name": "甲"}]},
                    {"dict_name": "D2", "dict_code": "DY", "items": []},
                ],
            },
        )
        assert r.status_code == 201
        std_id = r.json()["standard_id"]
        r2 = await client.post(f"/api/v1/data-standards-v2/{std_id}/export?format=excel",
                               headers={"X-API-Key": raw})
        assert r2.status_code == 200
        assert r2.content[:2] == b"PK"
        assert r2.headers["content-type"].startswith(
            "application/vnd.openxmlformats")

    async def test_export_not_found(self, client, db_session):
        raw, _ = await _key(db_session, "ie_x1_b")
        r = await client.post(
            f"/api/v1/data-standards-v2/{uuid.uuid4()}/export?format=json",
            headers={"X-API-Key": raw},
        )
        assert r.status_code == 404

    async def test_export_foreign_owner(self, client, db_session):
        raw_a, ka = await _key(db_session, "ie_x1_c")
        raw_b, _ = await _key(db_session, "ie_x1_c2")
        r = await client.post("/api/v1/data-standards-v2/import/json",
                              headers={"X-API-Key": raw_a},
                              json={"name": "越权标准"})
        std_id = r.json()["standard_id"]
        r2 = await client.post(f"/api/v1/data-standards-v2/{std_id}/export?format=json",
                               headers={"X-API-Key": raw_b})
        assert r2.status_code == 404

    async def test_export_bad_format(self, client, db_session):
        """真正不支持的格式仍 400。

        v2.22.2 此处曾以 `csv` 作反例（当时后端只实现 json/excel）；
        v2.22.3 补齐 csv 导出分支后，反例改为 `xml`。
        """
        raw, _ = await _key(db_session, "ie_x1_d")
        r = await client.post("/api/v1/data-standards-v2/import/json",
                              headers={"X-API-Key": raw},
                              json={"name": "格式标准"})
        std_id = r.json()["standard_id"]
        r2 = await client.post(f"/api/v1/data-standards-v2/{std_id}/export?format=xml",
                               headers={"X-API-Key": raw})
        assert r2.status_code == 400
        assert "不支持的导出格式" in r2.json()["detail"]


# ═══════════════════════════════════════════════════════════════════════════
# auto-create-table
# ═══════════════════════════════════════════════════════════════════════════

class TestAutoCreateTable:

    async def _seed_std_and_files(self, db_session, key, with_schema=True):
        std = DataStandard(id=str(uuid.uuid4()), api_key_id=key.id,
                           name="自动建表标准", code="AC-1")
        db_session.add(std)
        await db_session.flush()
        files = []
        if with_schema:
            sf = StoredFile(id=str(uuid.uuid4()), api_key_id=key.id,
                            original_name="地块.csv", file_data=b"x",
                            field_schema=[{"name": "id", "type": "string"},
                                          {"name": "geom", "type": "geometry"}])
            db_session.add(sf)
            files.append(sf)
        return std, files

    async def test_auto_create_ok(self, client, db_session):
        raw, key = await _key(db_session, "ie_au_a")
        std, [sf] = await self._seed_std_and_files(db_session, key)
        await db_session.commit()
        r = await client.post(
            f"/api/v1/data-standards-v2/{std.id}/auto-create-table",
            headers={"X-API-Key": raw},
            data={"file_id": sf.id, "table_name": "测试表", "geometry_type": "Point"},
        )
        assert r.status_code == 200
        body = r.json()
        assert body["field_count"] == 2
        assert body["table_name"] == "测试表"
        assert body["layer_group_name"] == "地块"  # 默认取文件名 stem

    async def test_auto_create_file_not_found(self, client, db_session):
        raw, _ = await _key(db_session, "ie_au_b")
        r = await client.post(
            f"/api/v1/data-standards-v2/{uuid.uuid4()}/auto-create-table",
            headers={"X-API-Key": raw},
            data={"file_id": str(uuid.uuid4()), "table_name": "T"},
        )
        assert r.status_code == 404

    async def test_auto_create_no_schema(self, client, db_session):
        raw, key = await _key(db_session, "ie_au_c")
        std = DataStandard(id=str(uuid.uuid4()), api_key_id=key.id, name="无结构标准")
        db_session.add(std)
        sf = StoredFile(id=str(uuid.uuid4()), api_key_id=key.id,
                        original_name="x.csv", file_data=b"x", field_schema=None)
        db_session.add_all([std, sf])
        await db_session.commit()
        r = await client.post(
            f"/api/v1/data-standards-v2/{std.id}/auto-create-table",
            headers={"X-API-Key": raw},
            data={"file_id": sf.id, "table_name": "T"},
        )
        assert r.status_code == 400
        assert "field_schema" in r.json()["detail"]

    async def test_auto_create_value_error(self, client, db_session):
        raw, key = await _key(db_session, "ie_au_d")
        std, [sf] = await self._seed_std_and_files(db_session, key)
        await db_session.commit()
        with patch(
            "app.services.structure_analyzer.auto_create_standard_from_schema",
            new=AsyncMock(side_effect=ValueError("几何类型不受支持")),
        ):
            r = await client.post(
                f"/api/v1/data-standards-v2/{std.id}/auto-create-table",
                headers={"X-API-Key": raw},
                data={"file_id": sf.id, "table_name": "T"},
            )
        assert r.status_code == 400
        assert "几何类型不受支持" in r.json()["detail"]

    async def test_auto_create_foreign_owner(self, client, db_session):
        raw_a, key_a = await _key(db_session, "ie_au_e")
        raw_b, _ = await _key(db_session, "ie_au_e2")
        std = DataStandard(id=str(uuid.uuid4()), api_key_id=key_a.id, name="他人文件")
        db_session.add(std)
        await db_session.flush()
        sf = StoredFile(id=str(uuid.uuid4()), api_key_id=key_a.id,
                        original_name="a.csv", file_data=b"x",
                        field_schema=[{"name": "f", "type": "string"}])
        db_session.add(sf)
        await db_session.commit()
        r = await client.post(
            f"/api/v1/data-standards-v2/{std.id}/auto-create-table",
            headers={"X-API-Key": raw_b},
            data={"file_id": sf.id, "table_name": "T"},
        )
        assert r.status_code == 404


# ═══════════════════════════════════════════════════════════════════════════
# 实现层函数直调（route 不经过的边角）
# ═══════════════════════════════════════════════════════════════════════════

class TestHelpersDirect:

    async def test_import_version_snapshot_with_extra_diff(self, db_session):
        from app.core.auth_context import AuthContext
        raw, key = await _key(db_session, "ie_hl_a")
        std = DataStandard(id=str(uuid.uuid4()), api_key_id=key.id,
                           name="快照标准", version=1)
        db_session.add(std)
        await db_session.commit()
        ctx = AuthContext(api_key=key, accessible_key_ids=[key.id])
        v = await ie._import_version_snapshot(
            db_session, ctx, std, "直调日志",
            extra_diff={"added": {"tables": ["T1"]}},
        )
        assert v == 2
        rows = (await db_session.execute(
            select(StandardVersion).where(StandardVersion.standard_id == std.id)
        )).scalars().all()
        assert len(rows) == 1
        assert rows[0].version_number == 1
        assert rows[0].change_log == "直调日志"
        assert rows[0].diff_data["added"]["tables"] == ["T1"]

    def test_content_disposition_fallback(self):
        header = ie._content_disposition("纯中文导出")
        assert "filename=standard_export" in header
        assert "filename*=UTF-8''" in header

    def test_content_disposition_ascii(self):
        header = ie._content_disposition("report.json")
        assert 'filename=report.json' in header
        assert "UTF-8" in header